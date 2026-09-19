#!/usr/bin/env python3
"""
remove-dd-rows.py — remove resources, fields and lookups from a DD reference XLSX, keeping every other
row, its formatting and its hyperlinks intact.

Use when elements have to come out of a version's sheet (an element entered ahead of its proposal's
ratification, a retraction voted by the workgroup). Rows are removed from the Fields, Lookups and
Changes tabs together, so the changelog does not keep claiming the version added them.

Usage
-----
    python3 remove-dd-rows.py <input.xlsx> <output.xlsx> \
        [--resource Model ...] \
        [--field Field.CollectionYN Lookup.FeedTypes ...] \
        [--lookup ModelType FieldDataTypes ...]

    --resource NAME        drop the resource: every Fields row with that ResourceName, and every Changes
                           row for the resource or its fields
    --field RES.NAME       drop one field (Fields row + its Changes rows)
    --lookup NAME          drop a whole lookup: every Lookups row with that LookupName + its Changes rows

openpyxl only (the sanctioned XLSX writer). Its delete_rows moves values and styles but not
hyperlinks: they stay bound to their old coordinates, and on the next load an orphaned hyperlink
materializes as a ghost cell whose value is the target URL. So each tab is processed as: capture the
hyperlinks of every row, delete the rows, clear every hyperlink, re-apply each surviving row's
hyperlinks at its new position, then drop the trailing cells the pass instantiated so the saved
dimension ends at the last real row (read-only readers, the linter among them, trust the dimension).
Run dd-sheet-linter.py, lint-dd-sheet.py and diff-dd-sheet.py on the output.
"""
import argparse
import sys

import openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink


def remove_rows(ws, predicate):
    """Delete every data row for which predicate(values, header_index) is true. Returns the removed rows."""
    header = [c.value for c in ws[1]]
    index = {h: i for i, h in enumerate(header)}
    ncol = len(header)
    last = ws.max_row
    links = {}
    victims = []
    for r in range(2, last + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, ncol + 1)]
        row_links = {}
        for c in range(1, ncol + 1):
            hl = ws.cell(row=r, column=c).hyperlink
            if hl is not None:
                row_links[c] = (hl.target, hl.display, hl.tooltip)
        if row_links:
            links[r] = row_links
        if predicate(values, index):
            victims.append((r, values))
    victim_rows = {r for r, _ in victims}
    for r, _ in sorted(victims, reverse=True):
        ws.delete_rows(r, 1)
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink = None
    new_r = 1
    for r in range(2, last + 1):
        if r in victim_rows:
            continue
        new_r += 1
        for c, (target, display, tooltip) in links.get(r, {}).items():
            cell = ws.cell(row=new_r, column=c)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=target, display=display, tooltip=tooltip)
    if victims:
        ws.delete_rows(new_r + 1, len(victims))
    return [v for _, v in victims]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--resource", nargs="*", default=[], metavar="NAME")
    ap.add_argument("--field", nargs="*", default=[], metavar="RES.NAME")
    ap.add_argument("--lookup", nargs="*", default=[], metavar="NAME")
    args = ap.parse_args()
    resources = set(args.resource)
    fields = set()
    for f in args.field:
        if "." not in f:
            sys.exit(f"--field expects ResourceName.StandardName, got {f!r}")
        fields.add(tuple(f.split(".", 1)))
    lookups = set(args.lookup)
    if not (resources or fields or lookups):
        sys.exit("nothing to remove")

    wb = openpyxl.load_workbook(args.input)

    def fields_pred(v, i):
        res, name = v[i["ResourceName"]], v[i["StandardName"]]
        return res in resources or (res, name) in fields

    def lookups_pred(v, i):
        return v[i["LookupName"]] in lookups

    def changes_pred(v, i):
        res, name, kind = v[i["ResourceName"]], v[i["FieldName"]], v[i["Data Element Type"]]
        if res in resources:
            return True
        if kind == "Field" and (res, name) in fields:
            return True
        if kind == "Lookup" and name in lookups:
            return True
        return False

    removed = {
        "Fields": remove_rows(wb["Fields"], fields_pred),
        "Lookups": remove_rows(wb["Lookups"], lookups_pred),
        "Changes": remove_rows(wb["Changes"], changes_pred) if "Changes" in wb.sheetnames else [],
    }
    wb.save(args.output)
    for tab, rows in removed.items():
        print(f"{tab}: {len(rows)} row(s) removed")
    print(f"Wrote: {args.output}")


if __name__ == "__main__":
    main()
